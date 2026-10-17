"""Real local HTTP/restart/MOCK SOC acceptance. Optional Docker engine uses same checks."""

import argparse
import json
import os
import secrets
import shutil
import socket
import subprocess
import sys
from io import BytesIO
from pathlib import Path
from time import perf_counter, sleep
from uuid import uuid4

import httpx
from openpyxl import load_workbook


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--docker-image")
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    directory = root / "outputs/completion-acceptance" / uuid4().hex
    directory.mkdir(parents=True)
    if args.docker_image and not shutil.which("docker"):
        result = {
            "status": "BLOCKED",
            "reason": "docker_engine_unavailable",
            "real_business": "BLOCKED",
        }
        (directory / "acceptance.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
        print(json.dumps({"report_directory": str(directory), **result}))
        return 2
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
    token = secrets.token_urlsafe(32)
    other = secrets.token_urlsafe(32)
    env = {
        **os.environ,
        "AGENT_DEMO_API_TOKENS": json.dumps({"operator": token, "other": other}),
        "AGENT_DEMO_DATA_ROOT": str(directory),
    }
    auth = {"Authorization": "Bearer " + token}
    second = {"Authorization": "Bearer " + other}
    child = None
    name = "agent-demo-check-" + uuid4().hex[:12]
    container_created = False
    logs = (directory / "server.log").open("wb")

    def docker(*arguments):
        return subprocess.run(
            ["docker", *arguments], env=env, capture_output=True, check=True, timeout=90
        )

    def start():
        nonlocal child, container_created
        if args.docker_image:
            if container_created:
                docker("start", name)
            else:
                docker(
                    "run",
                    "-d",
                    "--name",
                    name,
                    "-p",
                    f"127.0.0.1:{port}:8080",
                    "--mount",
                    f"type=volume,source={name}-runtime,target=/app/runtime",
                    "--mount",
                    f"type=volume,source={name}-outputs,target=/app/outputs",
                    "-e",
                    "AGENT_DEMO_API_TOKENS",
                    args.docker_image,
                )
                container_created = True
        else:
            child = subprocess.Popen(
                [
                    sys.executable,
                    "-m",
                    "uvicorn",
                    "agent_demo.api.app:create_app",
                    "--factory",
                    "--host",
                    "127.0.0.1",
                    "--port",
                    str(port),
                    "--workers",
                    "1",
                ],
                cwd=root,
                env=env,
                stdout=logs,
                stderr=logs,
            )
        deadline = perf_counter() + 30
        while perf_counter() < deadline:
            try:
                if (
                    httpx.get(
                        f"http://127.0.0.1:{port}/health", timeout=1, trust_env=False
                    ).status_code
                    == 200
                ):
                    return
            except httpx.HTTPError:
                pass
            sleep(0.1)
        raise RuntimeError("local_server_start_failed")

    def stop():
        nonlocal child
        if args.docker_image:
            if container_created:
                docker("stop", "-t", "15", name)
        elif child is not None:
            child.terminate()
            child.wait(timeout=20)
            child = None

    checks = []

    def check(name, condition):
        checks.append({"check": name, "passed": bool(condition)})
        if not condition:
            raise AssertionError(name)

    try:
        start()
        payload = json.loads((root / "data/fixtures/soc_requirements.json").read_text("utf-8"))
        payload.pop("evidences")  # API only permits trusted server-side evidence paths.
        with httpx.Client(
            base_url=f"http://127.0.0.1:{port}", headers=auth, timeout=30, trust_env=False
        ) as client:

            def wait(task):
                until = perf_counter() + 20
                while perf_counter() < until:
                    response = client.get("/tasks/" + task)
                    response.raise_for_status()
                    state = response.json()
                    if state["status"] not in ("pending", "running"):
                        return state
                    sleep(0.05)
                raise TimeoutError("task_wait_timeout")

            def approval(state, action="approve", **extra):
                snapshot = state["review_snapshot"]
                return dict(
                    action=action,
                    review_id=snapshot["review_id"],
                    review_version=snapshot["review_version"],
                    snapshot_digest=snapshot["snapshot_digest"],
                    **extra,
                )

            result = client.post("/tasks", json=payload)
            check("submit_202", result.status_code == 202)
            task = result.json()["task_id"]
            state = wait(task)
            check("awaiting_review", state["status"] == "awaiting_review")
            check(
                "unapproved_export_denied",
                client.get("/tasks/" + task + "/result").status_code == 409,
            )
            check(
                "other_owner_denied",
                client.get("/tasks/" + task, headers=second).status_code == 404,
            )
            stop()
            start()
            restored = client.get("/tasks/" + task).json()
            check("restart_same_review", restored["review_snapshot"] == state["review_snapshot"])
            requirements = payload["requirements"]
            requirements[0]["expected_value"] = 96
            modified = client.post(
                "/tasks/" + task + "/review",
                json=approval(restored, "modify", modification={"requirements": requirements}),
            )
            check("modify_accepted", modified.status_code == 200)
            newer = modified.json()
            check("new_review_version", newer["review_snapshot"]["review_version"] == 2)
            check(
                "stale_approval_denied",
                client.post("/tasks/" + task + "/review", json=approval(restored)).status_code
                == 409,
            )
            action = approval(newer)
            approved = client.post("/tasks/" + task + "/review", json=action)
            check(
                "approved_completed",
                approved.status_code == 200 and approved.json()["status"] == "completed",
            )
            first = client.get("/tasks/" + task + "/result")
            check("xlsx_download", first.status_code == 200)
            check(
                "duplicate_approval",
                client.post("/tasks/" + task + "/review", json=action).status_code == 200,
            )
            repeat = client.get("/tasks/" + task + "/result")
            check("idempotent_bytes", first.content == repeat.content)
            stop()
            start()
            check(
                "completed_restart_export",
                client.get("/tasks/" + task + "/result").content == first.content,
            )
            book = load_workbook(BytesIO(first.content))
            rows = list(book["SOC"].iter_rows(min_row=4, values_only=True))
            expected = json.loads((root / "data/fixtures/soc_expected_mock.json").read_text())[
                "expected_statuses"
            ]
            expected["r1"] = "NC"
            check("16_rows_statuses", len(rows) == 16 and {r[0]: r[5] for r in rows} == expected)
            check("literal_cells", all(c.data_type != "f" for s in book for row in s for c in row))
            check(
                "approved_draft",
                dict(book["Review"].values)["human_review_status"] == "approved"
                and dict(book["Review"].values)["document_status"] == "DRAFT",
            )
            book.close()
            (directory / "MOCK-SOC-DRAFT.xlsx").write_bytes(first.content)
            task2 = client.post("/tasks", json=payload).json()["task_id"]
            second_state = wait(task2)
            rejected = client.post(
                "/tasks/" + task2 + "/review", json=approval(second_state, "reject")
            )
            check(
                "reject_cancelled",
                rejected.status_code == 200 and rejected.json()["status"] == "cancelled",
            )
            check(
                "reject_export_denied", client.get("/tasks/" + task2 + "/result").status_code == 409
            )
        result = {
            "status": "PASS",
            "mode": "docker_http_mock" if args.docker_image else "local_http_mock",
            "checks": checks,
            "task_id": task,
            "review_version": 2,
            "real_business": "BLOCKED",
            "limitations": ["MOCK evidence only", "not formal customer publication"],
        }
        code = 0
    except Exception as error:
        result = {
            "status": "FAIL",
            "checks": checks,
            "error_type": type(error).__name__,
            "real_business": "BLOCKED",
        }
        code = 1
    finally:
        stop()
        if args.docker_image and container_created:
            docker("rm", name)  # Only this invocation's UUID-named disposable container.
            docker("volume", "rm", name + "-runtime", name + "-outputs")
        logs.close()
    (directory / "acceptance.json").write_text(
        json.dumps(result, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps({"report_directory": str(directory), **result}))
    return code


if __name__ == "__main__":
    raise SystemExit(main())
