"""Deterministic XLSX drafts; publication uses the Phase 5 operation ledger."""

from datetime import datetime
from io import BytesIO
from math import ceil
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from agent_demo.domain.decisions import WorkflowResult
from agent_demo.graphs.review_graph import digest, snapshot_digest

HEADERS = [
    "Item No.",
    "Parent ID",
    "Original Requirement",
    "Sub Requirement",
    "Proposed Huawei Model",
    "Compliance Status",
    "Compliance Explanation",
    "Evidence IDs",
    "Evidence Extract",
    "Source Document",
    "Source Version",
    "Source Page / Chunk",
    "Human Review",
    "Review Version",
    "Requirement Source Location",
]


def write_cell(sheet, row, col, value):
    if value is None:
        value = "Unknown"
    if isinstance(value, str) and len(value) > 32767:
        raise ValueError("Excel cell length exceeded; split evidence before export")
    cell = sheet.cell(row, col, value)
    if isinstance(value, str):
        # Untrusted requirements and evidence are literal text, never formulas.
        cell.data_type = "s"
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    cell.font = Font(name="Arial", size=11)
    return cell


def workbook_bytes(state, audit):
    snapshot = state["review_snapshot"]
    if snapshot_digest(snapshot) != snapshot["snapshot_digest"]:
        raise ValueError("snapshot changed")
    business = WorkflowResult.model_validate(snapshot["business_result"])
    if state["status"] in {"rejected", "failed", "cancelled"}:
        raise ValueError("this task cannot be exported")
    reviewed = state["status"] == "completed"
    if reviewed and (not state.get("approval") or audit.get("action") != "approve"):
        raise ValueError("approval and application audit required")
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "AGENT-DEMO"
    wb.properties.created = wb.properties.modified = datetime(2000, 1, 1)
    soc, evidence, review = [wb.create_sheet(n) for n in ("SOC", "Evidence", "Review")]
    for sheet in wb:
        sheet.sheet_view.showGridLines = False
    soc.append(["SOC DRAFT"])
    soc.append(
        [f"{business.source_type} | human review: " + ("approved" if reviewed else "not_reviewed")]
    )
    soc.append(HEADERS)
    reqs = {r.id: r for r in business.requirements}
    evs = {e.id: e for e in business.evidences}
    decisions = {d.requirement_id: d for d in business.decisions}
    if set(reqs) != set(decisions):
        raise ValueError("missing or unexpected decisions")
    for i, req in enumerate(business.requirements, 4):
        decision = decisions[req.id]
        if any(eid not in evs for eid in decision.evidence_ids):
            raise ValueError("dangling evidence ID")
        selected = [evs[eid] for eid in decision.evidence_ids]
        original = reqs[req.parent_id].original_text if req.parent_id else req.original_text
        row = [
            req.id,
            req.parent_id,
            original,
            (
                f"{req.attribute} {req.operator} {req.expected_value} {req.unit}; "
                f"conditions: {req.applicable_conditions}"
                if req.kind == "atomic"
                else "Composite: see child requirement rows"
            ),
            req.model,
            decision.proposed_status,
            "\n".join(
                [decision.reason]
                + (
                    ["Missing: " + "; ".join(decision.missing_items)]
                    if decision.missing_items
                    else []
                )
                + (
                    ["Conflicts: " + "; ".join(decision.conflict_items)]
                    if decision.conflict_items
                    else []
                )
            ),
            "\n".join(decision.evidence_ids),
            "\n\n".join(e.source_text or "Unknown" for e in selected),
            "\n".join(e.source_document or "Unknown" for e in selected),
            "\n".join(e.source_version or "Unknown" for e in selected),
            "\n".join(
                f"{e.page_number_semantics}: {e.page}; chunk: {e.chunk_id}" for e in selected
            ),
            "approved" if reviewed else "not_reviewed",
            snapshot["review_version"],
            req.source_location,
        ]
        for col, value in enumerate(row, 1):
            write_cell(soc, i, col, value)
        if selected:
            evidence_row = list(evs).index(selected[0].id) + 2
            soc.cell(i, 8).hyperlink = f"#Evidence!A{evidence_row}"
    evidence.append(
        [
            "Evidence ID",
            "Model",
            "Attribute",
            "Actual",
            "Unit",
            "Original Text",
            "Document",
            "Version",
            "Page",
            "Page Semantics",
            "Chunk ID",
            "Source Type",
        ]
    )
    for i, item in enumerate(business.evidences, 2):
        row = [
            item.id,
            item.model,
            item.attribute,
            item.actual_value,
            item.unit,
            item.source_text,
            item.source_document,
            item.source_version,
            item.page,
            item.page_number_semantics,
            item.chunk_id,
            item.source_type,
        ]
        for col, value in enumerate(row, 1):
            write_cell(evidence, i, col, value)
    review.append(["Field", "Value"])
    metadata = {
        "document_status": "DRAFT",
        "run_id": state["run_id"],
        "thread_id": state["thread_id"],
        "review_id": snapshot["review_id"],
        "review_version": snapshot["review_version"],
        "snapshot_digest": snapshot["snapshot_digest"],
        "source_type": business.source_type,
        "human_review_status": "approved" if reviewed else "not_reviewed",
        **audit,
    }
    for i, (key, value) in enumerate(metadata.items(), 2):
        write_cell(review, i, 1, key)
        write_cell(review, i, 2, value)
    for sheet, header, widths in (
        (soc, 3, [20, 20, 48, 48, 26, 20, 52, 30, 64, 42, 24, 35, 22, 20, 60]),
        (evidence, 1, [30, 26, 24, 20, 16, 64, 42, 24, 15, 30, 20, 20]),
        (review, 1, [28, 80]),
    ):
        for i, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(i)].width = width
        for cell in sheet[header]:
            cell.fill = PatternFill("solid", fgColor="26384A")
            cell.font = Font(name="Arial", color="FFFFFF", bold=True, size=11)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[header].height = 36
        for row in sheet.iter_rows(min_row=header + 1):
            lines = max(
                sum(
                    max(
                        1,
                        ceil(
                            sum(2 if ord(ch) > 127 else 1 for ch in part)
                            / max(1, widths[cell.column - 1] - 2)
                        ),
                    )
                    for part in str(cell.value or "").split("\n")
                )
                for cell in row
            )
            if lines * 15 + 8 > 409:
                raise ValueError("text exceeds readable Excel row height; split the input")
            sheet.row_dimensions[row[0].row].height = max(30, lines * 15 + 8)
        sheet.freeze_panes = f"C{header + 1}" if sheet is soc else f"A{header + 1}"
        sheet.auto_filter.ref = f"A{header}:{get_column_letter(len(widths))}{sheet.max_row}"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.print_title_rows = f"1:{header}"
    # Stable ZIP metadata and core modified timestamp make retries byte-identical.
    raw = BytesIO()
    wb.save(raw)
    output = BytesIO()
    with ZipFile(raw) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            content = source.read(name)
            if name == "docProps/core.xml":
                import re

                content = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]+", rb"\g<1>2000-01-01T00:00:00Z", content
                )
            info = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, content)
    content = output.getvalue()
    # Verify actual serialized IDs, row count, conclusions and literal cell types.
    with BytesIO(content) as stream:
        check = load_workbook(stream)
        rows = list(check["SOC"].iter_rows(min_row=4, values_only=True))
        assert [r[0] for r in rows] == list(reqs)
        assert [r[5] for r in rows] == [decisions[r].proposed_status for r in reqs]
        assert all(c.data_type != "f" for s in check for row in s for c in row)
        check.close()
    return content


def export_soc(state, audit, ledger):
    policy = state["review_snapshot"].get("export_policy") or {}
    if (
        policy.get("format_version") != "soc-v2"
        or policy.get("output_root") != str(ledger.output_root)
        or policy.get("document_status") != "DRAFT"
        or policy.get("kind") != "write_soc_excel_draft"
    ):
        raise ValueError("export type, directory and template must be included in human review")
    content = workbook_bytes(state, audit)
    operation = digest(
        {
            "export_version": "soc-v2",
            "snapshot": state["review_snapshot"],
            "audit": audit,
            "status": state["status"],
        }
    )
    return ledger.execute_bytes(operation, content, ledger.output_root / f"{operation}.xlsx")
