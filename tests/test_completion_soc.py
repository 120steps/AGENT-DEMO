"""Actual graph/SQLite/XLSX/PDF/API integration with isolated MOCK inputs."""

import base64
import json
import subprocess
import sys
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from threading import Event, Thread
from time import perf_counter, sleep

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from agent_demo.api.app import create_app
from agent_demo.domain.mock_cases import complete_evidence, requirements
from agent_demo.evaluation.fixtures import test_pdf as make_test_pdf
from agent_demo.integrations.rag_client import RAGClient, RAGSettings
from agent_demo.persistence.operations import OperationLedger
from agent_demo.soc.rfp import confirm_requirements, parse_rfp
from agent_demo.soc.service import SOCService, TaskNotFound

ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "data/fixtures/soc_requirements.json"
TOKENS = {"alice": "test-alice-token-0123456789abcdef", "bob": "test-bob-token-0123456789abcdef"}


def small_input():
    return dict(
        source_type="mock",
        requirements=requirements(),
        evidences=complete_evidence(),
        input_confirmed=True,
    )


def response(status, action="approve", **extra):
    snapshot = status["review_snapshot"]
    return dict(
        action=action,
        review_id=snapshot["review_id"],
        review_version=snapshot["review_version"],
        snapshot_digest=snapshot["snapshot_digest"],
        **extra,
    )


def ready(service, payload=None, owner="alice"):
    task = service.submit(payload or small_input(), owner)
    state = service.run(task, owner)
    assert state["status"] == "awaiting_review", state
    return task, state


def test_all_16_mock_rows_match_expected_and_xlsx_round_trip(tmp_path):
    service = SOCService(tmp_path)
    task, status = ready(service, json.loads(FIXTURE.read_text("utf-8")))
    expected = json.loads((FIXTURE.parent / "soc_expected_mock.json").read_text())[
        "expected_statuses"
    ]
    actual = {
        d["requirement_id"]: d["proposed_status"] for d in status["business_result"]["decisions"]
    }
    assert actual == expected
    with pytest.raises(ValueError):
        service.export(task, "alice")
    draft = service.export(task, "alice", draft=True)
    approved = service.review(task, "alice", response(status))
    assert approved["status"] == "completed"
    artifact = service.export(task, "alice")
    assert artifact["path"] != draft["path"]
    data = Path(artifact["path"]).read_bytes()
    book = load_workbook(BytesIO(data))
    rows = list(book["SOC"].iter_rows(min_row=4, values_only=True))
    assert len(rows) == 16 and {r[0]: r[5] for r in rows} == expected
    assert any("\u8bbe\u5907" in r[2] for r in rows)
    assert all(r[12] == "approved" for r in rows)
    assert dict(book["Review"].values)["reviewer"] == "alice"
    assert all(c.data_type != "f" for s in book for row in s for c in row)
    assert book["SOC"].freeze_panes == "C4"
    book.close()
    stamp = Path(artifact["path"]).stat().st_mtime_ns
    service.review(task, "alice", response(status))
    repeated = service.export(task, "alice")
    assert repeated["reused"] and Path(artifact["path"]).read_bytes() == data
    assert Path(artifact["path"]).stat().st_mtime_ns == stamp
    assert len(list((tmp_path / "outputs/soc").glob("*.xlsx"))) == 2


def test_reject_cancel_and_wrong_owner_have_no_protected_artifacts(tmp_path):
    service = SOCService(tmp_path)
    task, status = ready(service)
    for op in (service.status, service.run, service.cancel, service.export):
        with pytest.raises(TaskNotFound):
            op(task, "bob")
    service.review(task, "alice", response(status, "reject"))
    with pytest.raises(ValueError):
        service.export(task, "alice", draft=True)
    assert not list((tmp_path / "outputs").rglob("*.json"))
    assert not list((tmp_path / "outputs").rglob("*.xlsx"))
    task, status = ready(service)
    service.cancel(task, "alice")
    with pytest.raises(ValueError):
        service.review(task, "alice", response(status))
    assert service.run(task, "alice")["status"] == "cancelled"


def test_modify_rechecks_and_old_version_cannot_approve(tmp_path):
    service = SOCService(tmp_path)
    task, status = ready(service)
    changed = deepcopy(small_input()["requirements"])
    changed[1]["expected_value"] = 96
    newer = service.review(
        task, "alice", response(status, "modify", modification={"requirements": changed})
    )
    assert newer["review_snapshot"]["review_version"] == 2
    assert newer["business_result"]["overall_status"] == "NC"
    with pytest.raises(ValueError):
        service.review(task, "alice", response(status))
    service.review(task, "alice", response(newer))
    assert service.export(task, "alice")["status"] == "completed"


def test_untrusted_formula_is_literal_and_modified_approved_file_rejected(tmp_path):
    service = SOCService(tmp_path)
    payload = small_input()
    payload["requirements"][0]["original_text"] = '=HYPERLINK("file:///private","bad")'
    task, status = ready(service, payload)
    service.review(task, "alice", response(status))
    artifact = service.export(task, "alice")
    book = load_workbook(artifact["path"])
    assert book["SOC"]["C4"].data_type == "s" and book["SOC"]["C4"].value.startswith("=")
    book.close()
    Path(artifact["path"]).write_bytes(b"tampered")
    with pytest.raises(ValueError):
        service.export(task, "alice")


def test_sqlite_soc_reopens_across_real_python_processes(tmp_path):
    code = """
import json,sys
from agent_demo.soc.service import SOCService
from pathlib import Path
s=SOCService(Path(sys.argv[1]))
if sys.argv[2]=='start':
 p=json.loads(Path(sys.argv[3]).read_text('utf-8'))
 t=s.submit(p,'alice'); print(json.dumps(s.run(t,'alice')))
else:
 t=sys.argv[3]; state=s.status(t,'alice'); v=state['review_snapshot']
 s.review(t,'alice',dict(action='approve',review_id=v['review_id'],review_version=v['review_version'],snapshot_digest=v['snapshot_digest']))
 print(json.dumps(s.export(t,'alice')))
"""

    def process(action, arg):
        result = subprocess.run(
            [sys.executable, "-c", code, str(tmp_path), action, str(arg)],
            capture_output=True,
            text=True,
            timeout=30,
            check=True,
        )
        return json.loads(result.stdout)

    first = process("start", FIXTURE)
    assert first["status"] == "awaiting_review"
    second = process("resume", first["task_id"])
    assert Path(second["path"]).is_file()
    assert process("resume", first["task_id"])["reused"]


def test_binary_ledger_recovers_after_actual_process_exit(tmp_path):
    root = tmp_path / "outputs/soc"
    code = """
import os,sys
from pathlib import Path
from agent_demo.persistence.operations import OperationLedger
p=Path(sys.argv[1])
with OperationLedger(p/'runtime/binary.sqlite',p/'outputs/soc',artifact_kind='soc_excel') as l:
 l.execute_bytes('crash',b'controlled TEST bytes',p/'outputs/soc/crash.xlsx',
                 after_publish=lambda:os._exit(17))
"""
    crashed = subprocess.run([sys.executable, "-c", code, str(tmp_path)], timeout=20)
    assert crashed.returncode == 17
    stamp = (root / "crash.xlsx").stat().st_mtime_ns
    with OperationLedger(
        tmp_path / "runtime/binary.sqlite", root, artifact_kind="soc_excel"
    ) as ledger:
        artifact = ledger.execute_bytes("crash", b"controlled TEST bytes", "crash.xlsx")
        assert artifact["reused"] and ledger.get("crash")["status"] == "completed"
        with pytest.raises(ValueError):
            ledger.execute_bytes("escape", b"x", "../escape.xlsx")
    assert (root / "crash.xlsx").stat().st_mtime_ns == stamp


def test_real_path_is_mock_http_and_test_pdf_not_real_product(tmp_path):
    pdf = tmp_path / "test.pdf"
    sha = make_test_pdf(pdf)
    seen = []

    def handler(request):
        seen.append(json.loads(request.content))
        return httpx.Response(
            200,
            json={
                "documents": [
                    {
                        "document": "Downlink ports: 48 count",
                        "source": "TEST.pdf",
                        "page": 1,
                        "chunk_id": 1,
                        "document_id": "test-doc",
                        "version_id": "test-v1",
                    }
                ]
            },
        )

    settings = RAGSettings(
        base_url="http://127.0.0.1:8011",
        bearer_token="TEST-token",
        page_number_semantics="pdf_physical_1_based",
    )
    binding = dict(
        model="TEST-SW-48",
        document_id="test-doc",
        document_version="test-v1",
        sha256=sha,
        pdf_path=pdf,
        applicability_confirmed=True,
    )
    service = SOCService(
        tmp_path,
        rag_factory=lambda: RAGClient(settings, transport=httpx.MockTransport(handler)),
        bindings=[binding],
        allowed_pdf_root=tmp_path,
    )
    req = requirements()[1]
    req.update(model="TEST-SW-48", parent_id=None)
    task, status = ready(
        service, dict(source_type="real_rag", requirements=[req], input_confirmed=True)
    )
    assert seen and seen[0]["document_id"] == "test-doc"
    assert status["business_result"]["overall_status"] == "FC"
    service.review(task, "alice", response(status))
    assert service.export(task, "alice")["status"] == "completed"
    with pytest.raises(ValueError):
        service.submit(
            dict(source_type="real_rag", requirements=[req], evidences=[], input_confirmed=True),
            "alice",
        )


@pytest.mark.parametrize("failure", [False, True])
def test_real_path_empty_or_failure_never_becomes_nc(tmp_path, failure):
    pdf = tmp_path / "test.pdf"
    binding = dict(
        model="TEST-SW-48",
        document_id="d",
        document_version="v",
        sha256=make_test_pdf(pdf),
        pdf_path=pdf,
        applicability_confirmed=True,
    )
    settings = RAGSettings(base_url="http://127.0.0.1", bearer_token="TEST", max_retries=0)
    transport = httpx.MockTransport(
        lambda r: httpx.Response(503 if failure else 200, json={"documents": []})
    )
    service = SOCService(
        tmp_path,
        rag_factory=lambda: RAGClient(settings, transport=transport),
        bindings=[binding],
        allowed_pdf_root=tmp_path,
    )
    req = requirements()[1]
    req.update(model="TEST-SW-48", parent_id=None)
    task = service.submit(
        dict(source_type="real_rag", requirements=[req], input_confirmed=True), "alice"
    )
    result = service.run(task, "alice")
    assert result["business_result"]["overall_status"] == ("ERROR" if failure else "PENDING")
    assert result["status"] == ("failed" if failure else "awaiting_review")


def test_cancellation_during_retrieval_prevents_next_step(tmp_path, monkeypatch):
    service = SOCService(tmp_path)
    entered, release = Event(), Event()

    def retrieve(*args, **kwargs):
        entered.set()
        release.wait(3)
        return [], {}, []

    monkeypatch.setattr(service, "retrieve", retrieve)
    req = requirements()[1]
    req.update(model="TEST-SW", parent_id=None)
    task = service.submit(
        dict(source_type="real_rag", requirements=[req], input_confirmed=True), "alice"
    )
    worker = Thread(target=service.run, args=(task, "alice"))
    worker.start()
    assert entered.wait(2)
    service.cancel(task, "alice")
    release.set()
    worker.join(4)
    assert not worker.is_alive() and service.status(task, "alice")["status"] == "cancelled"
    assert service.status(task, "alice")["review_snapshot"] is None


def test_pdf_number_location_manual_mapping_and_omissions(tmp_path):
    p = tmp_path / "rfp.pdf"
    make_test_pdf(p, "1. At least 48 GE ports")
    parsed = parse_rfp(p.read_bytes())
    clause = parsed.clauses[0]
    assert clause.page == 1 and clause.number == "1."
    req = requirements()[1]
    req.update(
        parent_id=None,
        original_text=clause.original_text,
        source_location=f"sha256:{parsed.source_hash};page:1;clause:{clause.id}",
    )
    mapping = [dict(clause_id=clause.id, requirements=[req])]
    assert confirm_requirements(parsed, "MOCK-SW-48", mapping, confirmed=True)[0]["id"] == req["id"]
    for arguments in ({"confirmed": False}, {"confirmed": True, "mappings": []}):
        with pytest.raises(ValueError):
            confirm_requirements(
                parsed,
                "MOCK-SW-48",
                arguments.get("mappings", mapping),
                confirmed=arguments["confirmed"],
            )
    req["source_location"] = "page:999"
    with pytest.raises(ValueError):
        confirm_requirements(parsed, "MOCK-SW-48", mapping, confirmed=True)


def test_scanned_pdf_rejected_and_all_text_retained(tmp_path):
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=600, height=800)
    content = BytesIO()
    writer.write(content)
    result = parse_rfp(content.getvalue())
    assert result.warnings and result.pages[0]["original_text"] == ""
    with pytest.raises(ValueError):
        confirm_requirements(result, "MOCK-SW-48", [], confirmed=True)


def wait_task(client, task, headers):
    deadline = perf_counter() + 10
    while perf_counter() < deadline:
        result = client.get("/tasks/" + task, headers=headers)
        assert result.status_code == 200, result.text
        if result.json()["status"] not in {"pending", "running"}:
            return result.json()
        sleep(0.01)
    raise AssertionError("API task did not finish")


def test_api_auth_owner_review_export_and_restart(tmp_path):
    auth = {"Authorization": "Bearer " + TOKENS["alice"]}
    bob = {"Authorization": "Bearer " + TOKENS["bob"]}
    service = SOCService(tmp_path)
    payload = small_input()
    payload.pop("evidences")
    with TestClient(create_app(service=service, tokens=TOKENS)) as client:
        assert client.get("/health").status_code == 200
        assert client.post("/tasks", json=payload).status_code == 401
        result = client.post("/tasks", json=payload, headers=auth)
        assert result.status_code == 202, result.text
        task = result.json()["task_id"]
        state = wait_task(client, task, auth)
        for suffix, method in (
            ("", client.get),
            ("/run", client.post),
            ("/cancel", client.post),
            ("/result", client.get),
        ):
            assert method("/tasks/" + task + suffix, headers=bob).status_code == 404
        assert client.get("/tasks/" + task + "/result", headers=auth).status_code == 409
    with TestClient(create_app(service=SOCService(tmp_path), tokens=TOKENS)) as client:
        state = client.get("/tasks/" + task, headers=auth).json()
        assert state["status"] == "awaiting_review"
        approved = client.post("/tasks/" + task + "/review", json=response(state), headers=auth)
        assert approved.status_code == 200, approved.text
        result = client.get("/tasks/" + task + "/result", headers=auth)
        assert result.status_code == 200 and result.content[:2] == b"PK"
        assert client.post("/tasks/" + task + "/cancel", headers=auth).status_code == 409


def test_api_invalid_payload_pdf_and_rate_limit(tmp_path):
    auth = {"Authorization": "Bearer " + TOKENS["alice"]}
    with TestClient(
        create_app(service=SOCService(tmp_path), tokens=TOKENS, requests_per_minute=3)
    ) as client:
        invalid = small_input()
        invalid["url"] = "http://arbitrary"
        result = client.post("/tasks", json=invalid, headers=auth)
        assert result.status_code == 422 and "arbitrary" not in result.text
        assert (
            client.post("/rfp/parse", json={"pdf_base64": "bad"}, headers=auth).status_code == 422
        )
        pdf = tmp_path / "p.pdf"
        make_test_pdf(pdf, "1. MOCK specification")
        result = client.post(
            "/rfp/parse",
            json={"pdf_base64": base64.b64encode(pdf.read_bytes()).decode()},
            headers=auth,
        )
        assert result.status_code == 200 and result.json()["clauses"][0]["page"] == 1
        assert client.get("/tasks/unknown", headers=auth).status_code == 429


def test_crash_after_approved_checkpoint_preserves_reviewer_audit(tmp_path):
    service = SOCService(tmp_path)
    task, state = ready(service)
    code = """
import json,os,sys
from agent_demo.soc.service import SOCService
s=SOCService(sys.argv[1]); t=sys.argv[2]; v=s.status(t,'alice')['review_snapshot']
s.sync=lambda *args,**kwargs:os._exit(19)
s.review(t,'alice',dict(action='approve',review_id=v['review_id'],
 review_version=v['review_version'],snapshot_digest=v['snapshot_digest']))
"""
    crashed = subprocess.run([sys.executable, "-c", code, str(tmp_path), task], timeout=20)
    assert crashed.returncode == 19
    recovered = SOCService(tmp_path)
    status = recovered.status(task, "alice")
    assert status["status"] == "completed" and status["audit"]["reviewer"] == "alice"
    assert status["audit"]["action"] == "approve"
    assert recovered.export(task, "alice")["status"] == "completed"
    # A conflicting repeat must not poison the persisted successful audit.
    with pytest.raises(ValueError):
        recovered.review(task, "alice", response(state, "reject"))
    assert recovered.status(task, "alice")["audit"]["action"] == "approve"


def test_pdf_through_confirmation_graph_review_and_excel(tmp_path):
    p = tmp_path / "rfp.pdf"
    make_test_pdf(p, "1. At least 48 GE ports")
    parsed = parse_rfp(p.read_bytes())
    clause = parsed.clauses[0]
    req = requirements()[1]
    req.update(
        parent_id=None,
        original_text=clause.original_text,
        source_location=f"sha256:{parsed.source_hash};page:1;clause:{clause.id}",
    )
    confirmed = confirm_requirements(
        parsed, "MOCK-SW-48", [dict(clause_id=clause.id, requirements=[req])], confirmed=True
    )
    service = SOCService(tmp_path)
    task, state = ready(
        service, dict(source_type="mock", requirements=confirmed, input_confirmed=True)
    )
    service.review(task, "alice", response(state))
    artifact = service.export(task, "alice")
    book = load_workbook(artifact["path"])
    assert book["SOC"]["C4"].value == clause.original_text
    assert book["SOC"]["F4"].value == "FC"
    book.close()
    events = [
        json.loads(line) for line in (service.runtime / f"{task}.jsonl").read_text().splitlines()
    ]
    assert all(e["run_id"] == task and e["thread_id"] == task for e in events)
    assert any(e["node_name"] == "validate_review" for e in events)
    assert any(e["node_name"] == "compare_requirement" for e in events)


def test_api_queue_limit_creates_no_orphan_task(tmp_path, monkeypatch):
    service = SOCService(tmp_path)
    entered, release = Event(), Event()

    def blocked(*args):
        entered.set()
        release.wait(5)

    monkeypatch.setattr(service, "run", blocked)
    payload = small_input()
    payload.pop("evidences")
    auth = {"Authorization": "Bearer " + TOKENS["alice"]}
    try:
        with TestClient(create_app(service=service, tokens=TOKENS)) as client:
            for _ in range(4):
                assert client.post("/tasks", json=payload, headers=auth).status_code == 202
            assert entered.wait(2)
            assert client.post("/tasks", json=payload, headers=auth).status_code == 429
            with service.db() as db:
                assert db.execute("SELECT COUNT(*) FROM soc_tasks").fetchone()[0] == 4
            release.set()
    finally:
        release.set()


def test_total_task_wait_stops_and_does_not_start_review(tmp_path, monkeypatch):
    service = SOCService(tmp_path, task_timeout=0.04)
    gate = Event()

    def slow(*args, **kwargs):
        gate.wait(3)
        return [], {}, []

    monkeypatch.setattr(service, "retrieve", slow)
    req = requirements()[1]
    req.update(model="TEST-SW", parent_id=None)
    task = service.submit(
        dict(source_type="real_rag", requirements=[req], input_confirmed=True), "alice"
    )
    started = perf_counter()
    try:
        status = service.run(task, "alice")
        assert status["status"] == "failed" and status["error_type"] == "TimeoutError"
        assert status["review_snapshot"] is None and perf_counter() - started < 1
    finally:
        gate.set()


def test_export_scope_and_template_are_part_of_review_snapshot(tmp_path):
    service = SOCService(tmp_path)
    task, state = ready(service)
    policy = state["review_snapshot"]["export_policy"]
    assert policy["format_version"] == "soc-v2" and policy["document_status"] == "DRAFT"
    assert Path(policy["output_root"]) == tmp_path / "outputs/soc"
    service.review(task, "alice", response(state))
    from agent_demo.soc.excel import export_soc

    with service.graph(task) as graph:
        from agent_demo.graphs.review_graph import thread_config

        approved = graph.get_state(thread_config(task)).values
        # A different approved-directory name under a second root still needs
        # a new review; permission to one SOC export does not authorize another.
        with OperationLedger(
            tmp_path / "runtime/other.sqlite",
            tmp_path / "other/outputs/soc",
            artifact_kind="soc_excel",
        ) as ledger:
            with pytest.raises(ValueError):
                export_soc(approved, service.get(task, "alice")["audit"], ledger)
